"""Converte os .md desta pasta em .docx (pandoc) e dá altura aos quadros de registro.

Uso (a partir desta pasta):  python build_docx.py
Saída: docx/doc1_*.docx, docx/doc2_*.docx, docx/doc3_*.docx
Suba os .docx ao Google Drive e abra como Google Docs.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

AQUI = Path(__file__).parent
SAIDA = AQUI / "docx"
ALTURA_CM = 3.5  # altura da linha vazia dos quadros de registro


def bordas(tabela):
    """Contorno e grade em todas as células (o pandoc gera tabelas sem bordas)."""
    tblPr = tabela._tbl.tblPr
    antigo = tblPr.find(qn("w:tblBorders"))
    if antigo is not None:
        tblPr.remove(antigo)
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "808080")
        b.append(e)
    tblPr.append(b)


def e_quadro_de_registro(tabela):
    """Quadro = tabela de 1 coluna e 2 linhas, com a 2ª linha vazia."""
    return (
        len(tabela.columns) == 1
        and len(tabela.rows) == 2
        and not tabela.rows[1].cells[0].text.strip()
    )


def converter(md):
    SAIDA.mkdir(exist_ok=True)
    saida = SAIDA / (md.stem + ".docx")
    subprocess.run(["pandoc", str(md), "-o", str(saida)], check=True)
    doc = Document(saida)
    n = 0
    for t in doc.tables:
        bordas(t)
        if e_quadro_de_registro(t):
            t.rows[1].height = Cm(ALTURA_CM)
            n += 1
    doc.save(saida)
    print(f"{saida.name}: {len(doc.tables)} tabelas, {n} quadros de registro")


for md in sorted(AQUI.glob("doc[0-9]_*.md")):
    converter(md)
